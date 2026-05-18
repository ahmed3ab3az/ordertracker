"""Excel parsing & validation for bulk order imports.

* Expects columns: ``name``, ``address``, ``price`` (case-insensitive).
* If ANY row is invalid, the entire file is rejected and a list of errors is returned.
* Atomicity at the DB layer is handled by :class:`services.orders.OrderService.import_bulk`.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from ..constants import MAX_EXCEL_ROWS
from .orders import OrderImportInput

REQUIRED_COLUMNS = ("name", "address", "price")


@dataclass
class RowError:
    row_number: int
    column: str
    reason: str


@dataclass
class ImportPreview:
    rows: list[OrderImportInput] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def _normalise_header(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _coerce_price(raw, row_number: int, errors: list[RowError]) -> Decimal | None:
    if raw is None or str(raw).strip() == "":
        errors.append(RowError(row_number, "price", "Price is required."))
        return None
    try:
        amount = Decimal(str(raw).replace(",", "").strip())
    except InvalidOperation:
        errors.append(RowError(row_number, "price", f"Could not parse price: {raw!r}"))
        return None
    if amount < 0:
        errors.append(RowError(row_number, "price", "Price must be non-negative."))
        return None
    return amount.quantize(Decimal("0.01"))


def parse_excel(path: str | Path) -> ImportPreview:
    """Parse and validate an Excel file. Never touches the DB."""
    preview = ImportPreview()
    path = Path(path)

    if not path.exists():
        preview.errors.append(RowError(0, "-", f"File not found: {path}"))
        return preview

    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        preview.errors.append(RowError(0, "-", f"Cannot open workbook: {exc}"))
        return preview

    ws = wb.active
    if ws is None:
        preview.errors.append(RowError(0, "-", "Workbook has no active worksheet."))
        return preview

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        preview.errors.append(RowError(0, "-", "Workbook is empty."))
        return preview

    headers = [_normalise_header(c) for c in header_row]
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        preview.errors.append(
            RowError(1, "header", f"Missing required columns: {', '.join(missing)}")
        )
        return preview

    name_idx = headers.index("name")
    address_idx = headers.index("address")
    price_idx = headers.index("price")

    for i, row in enumerate(rows_iter, start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue  # ignore blank rows

        if i - 1 > MAX_EXCEL_ROWS:
            preview.errors.append(
                RowError(i, "-", f"Exceeded max rows ({MAX_EXCEL_ROWS}).")
            )
            break

        name = (row[name_idx] or "").__str__().strip() if name_idx < len(row) else ""
        address = (row[address_idx] or "").__str__().strip() if address_idx < len(row) else ""
        price_raw = row[price_idx] if price_idx < len(row) else None

        if not name:
            preview.errors.append(RowError(i, "name", "Name is required."))
        if not address:
            preview.errors.append(RowError(i, "address", "Address is required."))
        price = _coerce_price(price_raw, i, preview.errors)

        if name and address and price is not None:
            preview.rows.append(OrderImportInput(name=name, address=address, price=price))

    if not preview.rows and not preview.errors:
        preview.errors.append(RowError(0, "-", "No data rows found in workbook."))

    return preview
